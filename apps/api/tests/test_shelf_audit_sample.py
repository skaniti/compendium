"""shelf_audit_sample: blind sample, key, scoring and the dry-run / apply split.

Synthetic data only; the embedder and the model are stubbed. The first half
needs no database; the second half runs on the test database when reachable.
"""

import csv
import random
from pathlib import Path

import numpy as np
import pytest

from backend.config.settings import settings
from backend.scripts import shelf_audit_sample as sa
from backend.services.shelf_service import ShelfConfig


def _vec(deg):
    a = np.deg2rad(deg)
    return np.array([np.cos(a), np.sin(a)])


def cand(cid, subtopic, kv="keep", cluster=None, category="home", **label):
    lab = {
        "subtopic": subtopic,
        "category": category,
        "knowledge_value": kv,
        "content_kind": "article_reference",
    }
    lab.update(label)
    return sa.Candidate(
        cid,
        f"Synthetic title {cid}",
        f"https://p{cid}.example.test/x",
        "p.example.test",
        cluster,
        lab,
    )


def shelved(cid, shelf="Shelf one", cluster=None, kv="keep"):
    return sa.Shelved(cand(cid, shelf, kv=kv, cluster=cluster), shelf, "home", 3)


# --- output path guard ------------------------------------------------------


def test_output_inside_repo_is_refused(tmp_path):
    inside = sa.repo_root() / "apps" / "api" / "sample.csv"
    with pytest.raises(ValueError, match="inside the repository"):
        sa.check_output_path(inside, overwrite=False)
    with pytest.raises(ValueError, match="inside the repository"):
        sa.check_output_path(sa.repo_root() / "x" / ".." / "sample.csv", overwrite=True)


def test_output_must_be_new_csv_or_xlsx_in_existing_dir(tmp_path):
    ok = sa.check_output_path(tmp_path / "s.csv", overwrite=False)
    assert ok == (tmp_path / "s.csv").resolve()
    (tmp_path / "s.csv").write_text("marks")
    with pytest.raises(ValueError, match="--overwrite"):
        sa.check_output_path(tmp_path / "s.csv", overwrite=False)
    assert sa.check_output_path(tmp_path / "s.csv", overwrite=True)
    with pytest.raises(ValueError, match="csv or .xlsx"):
        sa.check_output_path(tmp_path / "s.txt", overwrite=False)
    with pytest.raises(ValueError, match="does not exist"):
        sa.check_output_path(tmp_path / "nope" / "s.csv", overwrite=False)
    assert sa.default_key_path(tmp_path / "s.xlsx") == tmp_path / "s.key.csv"


def test_symlink_into_repo_is_refused(tmp_path):
    link = tmp_path / "link"
    link.symlink_to(sa.repo_root())
    with pytest.raises(ValueError, match="inside the repository"):
        sa.check_output_path(link / "sample.csv", overwrite=False)


# --- in-memory shelves ------------------------------------------------------


def fake_embed(texts):
    # "alpha*" -> 0 deg, "alpha near" -> 10 deg (attach), "mid" -> 45.6 deg vs alpha (ask band),
    # "beta*" -> 90 deg
    out = []
    for t in texts:
        if t.startswith("alpha near"):
            out.append(_vec(10))
        elif t.startswith("alpha"):
            out.append(_vec(0))
        elif t.startswith("mid"):
            out.append(_vec(45.6))
        else:
            out.append(_vec(90))
    return np.array(out)


def test_in_memory_build_attach_create_and_ask_band_counted():
    cands = [
        cand(1, "Alpha Shelf"),
        cand(2, "Alpha near shelf"),  # cos 0.985 -> attach
        cand(3, "Beta Shelf"),
        cand(4, "Mid Topic"),  # cos 0.70 to alpha -> ask band, declined in dry run -> own shelf
        cand(5, "Alpha Shelf"),  # same string -> memo
    ]
    pool, stats = sa.build_shelves_in_memory(
        cands, fake_embed, ShelfConfig(attach=0.80, ask=0.65)
    )
    names = {s.candidate.content_id: s.shelf_name for s in pool}
    assert names[1] == names[2] == names[5] == "Alpha Shelf"
    assert names[3] == "Beta Shelf" and names[4] == "Mid Topic"
    assert stats == {"ask_band": 1, "shelves": 3}
    sizes = {s.shelf_name: s.shelf_pages for s in pool}
    assert sizes == {"Alpha Shelf": 3, "Beta Shelf": 1, "Mid Topic": 1}


def test_dry_run_in_memory_makes_no_model_call(monkeypatch):
    def boom(*a, **k):
        raise AssertionError("model called in dry run")

    monkeypatch.setattr("backend.services.shelf_service.SameShelfAsker.__call__", boom)
    sa.build_shelves_in_memory([cand(1, "Alpha Shelf"), cand(2, "Mid Topic")], fake_embed)


def test_shelvable_filters_excluded_and_blank():
    cs = [
        cand(1, "Real shelf"),
        cand(2, "Utility", kv="exclude"),
        cand(3, "   "),
        cand(4, "Login", content_kind="auth_login"),
        cand(5, "Marginal shelf", kv="marginal"),
    ]
    assert [c.content_id for c in sa.shelvable(cs)] == [1, 5]


# --- sample -----------------------------------------------------------------


def _pool(n_keep=150, n_marginal=20):
    pool = [
        shelved(i, f"Shelf {i % 7}", cluster=f"Cluster {i % 5}" if i % 4 else None)
        for i in range(n_keep)
    ]
    pool += [shelved(1000 + i, "Side shelf", kv="marginal") for i in range(n_marginal)]
    return pool


def test_sample_is_100_distinct_keep_pages():
    sheet, key = sa.draw_sample(_pool())
    assert len(sheet) == len(key) == 100
    ids = [k[1] for k in key]
    assert len(set(ids)) == 100 and all(i < 1000 for i in ids)  # marginal pages never drawn
    assert [r[0] for r in sheet] == list(range(1, 101))
    assert all(r[5] == "" and r[6] == "" for r in sheet)


def test_sample_is_deterministic_per_seed_and_varies_across_seeds():
    a1, k1 = sa.draw_sample(_pool(), seed=7)
    a2, k2 = sa.draw_sample(_pool(), seed=7)
    assert a1 == a2 and k1 == k2
    a3, _ = sa.draw_sample(_pool(), seed=8)
    assert a3 != a1
    assert sa.draw_sample(list(reversed(_pool())), seed=7)[0] == a1  # pool order is irrelevant


def test_order_is_randomised_and_key_matches_columns():
    sheet, key = sa.draw_sample(_pool(), seed=3)
    by_item = {k[0]: k for k in key}
    orders = {k[2] for k in key}
    assert orders == {"shelf", "cluster"}
    shelf_in_a = sum(1 for k in key if k[2] == "shelf")
    assert 25 < shelf_in_a < 75  # not systematically one side
    for row in sheet:
        k = by_item[row[0]]
        shelf, cluster = sa.display_name(k[3]), sa.display_name(k[6])
        assert (row[3], row[4]) == ((shelf, cluster) if k[2] == "shelf" else (cluster, shelf))
        assert k[7] == "dry-run"


def test_pages_outside_clusters_get_the_neutral_placeholder():
    sheet, key = sa.draw_sample([shelved(1, "Only shelf", cluster=None)], n=5)
    assert len(sheet) == 1
    assert sa.NO_GROUP in (sheet[0][3], sheet[0][4]) and key[0][6] == ""
    assert "cluster" not in (sheet[0][3] + sheet[0][4]).lower()  # wording names neither side


def test_display_name_is_one_rule_for_both_columns():
    assert sa.display_name("whales and  cetaceans") == "Whales And Cetaceans"
    assert sa.display_name("Whales And Cetaceans") == "Whales And Cetaceans"
    assert sa.display_name("iPhone cases") == "IPhone Cases"
    assert sa.display_name(None) == sa.display_name("") == sa.display_name("   ") == sa.NO_GROUP


def test_casing_does_not_reveal_the_side():
    """Same words in different raw casing look identical in the sheet; the key keeps raw."""
    pool = [
        sa.Shelved(cand(i, "x", cluster="whales and cetaceans"), "Whales And Cetaceans", "c", 3)
        for i in range(60)
    ]
    sheet, key = sa.draw_sample(pool, n=60, seed=5)
    assert {(r[3], r[4]) for r in sheet} == {("Whales And Cetaceans", "Whales And Cetaceans")}
    assert {k[6] for k in key} == {"whales and cetaceans"}


def test_small_pool_returns_everything():
    sheet, _ = sa.draw_sample([shelved(i, "S") for i in range(4)], n=100)
    assert len(sheet) == 4
    assert sa.draw_sample([], n=100) == ([], [])


# --- files ------------------------------------------------------------------


def test_csv_roundtrip_keeps_unicode_and_commas(tmp_path):
    p = tmp_path / "s.csv"
    rows = [[1, "Synthetic, title é", "https://a.example.test/x?y=1,2", "A", "B", "", ""]]
    sa.write_table(p, sa.SHEET_HEADER, rows)
    header, back = sa.read_table(p)
    assert header == sa.SHEET_HEADER and back == [[str(v) for v in rows[0]]]


def test_xlsx_has_header_validation_and_roundtrips(tmp_path):
    from openpyxl import load_workbook

    p = tmp_path / "s.xlsx"
    sheet, _ = sa.draw_sample(_pool(20, 0), n=10)
    sa.write_table(p, sa.SHEET_HEADER, sheet, marks_cols=(5, 6))
    ws = load_workbook(p).active
    assert [c.value for c in ws[1]] == sa.SHEET_HEADER and ws.freeze_panes == "A2"
    ranges = {rng for r in ws.data_validations.dataValidation for rng in str(r.sqref).split()}
    assert ranges == {"F2:F11", "G2:G11"}
    header, back = sa.read_table(p)
    assert header == sa.SHEET_HEADER and len(back) == 10 and back[0][0] == "1"


# --- scoring ----------------------------------------------------------------


def _marked(sheet, key, shelf_right, cluster_right):
    """Fill marks: ``shelf_right``/``cluster_right`` are sets of 1-based items judged right."""
    by_item = {k[0]: k for k in key}
    out = []
    for row in sheet:
        k = by_item[row[0]]
        row = list(row)
        a_shelf = k[2] == "shelf"
        sa_ok, cl_ok = row[0] in shelf_right, row[0] in cluster_right
        row[5] = "right" if (sa_ok if a_shelf else cl_ok) else "wrong"
        row[6] = "right" if (cl_ok if a_shelf else sa_ok) else "wrong"
        out.append(row)
    return out


def test_score_attributes_marks_to_the_right_grouping():
    sheet, key = sa.draw_sample(_pool(), seed=1)
    shelf_right = set(range(1, 91))  # 90 of 100
    cluster_right = set(range(1, 61))  # 60 of 100
    res = sa.score_marks(_marked(sheet, key, shelf_right, cluster_right), key)
    assert (res["shelf_right"], res["shelf_marked"]) == (90, 100)
    assert (res["cluster_right"], res["cluster_marked"]) == (60, 100)
    assert res["shelf_rate"] == pytest.approx(0.9) and res["lift"] == pytest.approx(0.3)
    assert res["complete"] and res["accept"]


def test_score_rule_needs_both_rate_and_lift():
    sheet, key = sa.draw_sample(_pool(), seed=1)
    high_but_no_lift = sa.score_marks(
        _marked(sheet, key, set(range(1, 86)), set(range(1, 81))), key
    )
    assert high_but_no_lift["shelf_rate"] == 0.85 and not high_but_no_lift["accept"]
    lift_but_low = sa.score_marks(_marked(sheet, key, set(range(1, 71)), set(range(1, 31))), key)
    assert lift_but_low["lift"] == pytest.approx(0.4) and not lift_but_low["accept"]
    at_threshold = sa.score_marks(_marked(sheet, key, set(range(1, 81)), set(range(1, 61))), key)
    assert at_threshold["accept"]


def test_score_incomplete_is_never_accepted_and_accepts_aliases():
    sheet, key = sa.draw_sample(_pool(), seed=1, n=10)
    rows = _marked(sheet, key, set(range(1, 11)), set())
    rows[0][5] = rows[0][6] = ""
    res = sa.score_marks(rows, key)
    assert not res["complete"] and not res["accept"] and res["shelf_marked"] == 9
    assert "NOT COMPLETE" in sa.format_score(res)
    assert [sa._mark(v) for v in ("Right", " y ", "X", "WRONG", "n", "", None, "maybe")] == [
        "right",
        "right",
        "right",
        "wrong",
        "wrong",
        None,
        None,
        None,
    ]
    assert sa.score_marks([], key)["shelf_rate"] is None


def test_main_score_exit_codes(tmp_path, capsys):
    sheet, key = sa.draw_sample(_pool(), seed=2)
    kp = tmp_path / "s.key.csv"
    sa.write_table(kp, sa.KEY_HEADER, key)
    for name, shelf_n, cluster_n, code in (("good", 90, 50, 0), ("bad", 60, 50, 4)):
        p = tmp_path / f"{name}.csv"
        sa.write_table(
            p,
            sa.SHEET_HEADER,
            _marked(sheet, key, set(range(1, shelf_n + 1)), set(range(1, cluster_n + 1))),
        )
        assert sa.main(["--score", str(p), "--key", str(kp)]) == code
    out = capsys.readouterr().out
    assert "ACCEPT" in out and "NOT ACCEPTED" in out
    with pytest.raises(SystemExit):
        sa.main(["--score", str(tmp_path / "good.csv")])


def test_main_argument_errors(tmp_path, capsys):
    with pytest.raises(SystemExit):
        sa.main(["--user-id", "1"])  # no --out
    with pytest.raises(SystemExit):
        sa.main(["--user-id", "1", "--out", str(tmp_path / "s.csv"), "--maintain"])
    assert sa.main(["--user-id", "1", "--out", str(sa.repo_root() / "s.csv")]) == 2
    assert "inside the repository" in capsys.readouterr().err


# --- database ---------------------------------------------------------------

from tests.test_page_repo_time_windows import _pg_reachable

db = pytest.mark.skipif(not _pg_reachable(), reason="Test PostgreSQL not reachable")
if _pg_reachable():
    from tests import _label_fixtures as fx


@pytest.fixture
def world(monkeypatch):
    fx.sql(
        "TRUNCATE shelf_strings, shelf_pages, shelves, page_labels, cost_events, graph_cache,"
        " cluster_edges, page_clusters, clusters, recluster_runs, pages, page_content, captures,"
        " users CASCADE"
    )
    user = fx.make_user("auditor")
    cap = fx.make_capture(user)
    from backend.db import page_label_repo

    def add(name, subtopic, kv="keep", status="active", minute=0):
        pid, cid = fx.make_page(user, cap, name, minute=minute, status=status)
        page_label_repo.upsert_labels(
            [
                {
                    **fx.item(pid, subtopic=subtopic, knowledge_value=kv, category="home"),
                    "page_content_id": cid,
                    "user_id": user["id"],
                    "labeller_version": "synthetic@v",
                }
            ]
        )
        return pid, cid

    ids = {}
    for i in range(6):
        ids[f"a{i}"] = add(f"alpha{i}", "Alpha Shelf", minute=i)
    for i in range(4):
        ids[f"b{i}"] = add(f"beta{i}", "Beta Shelf", minute=10 + i)
    ids["x"] = add("util", "Login", kv="exclude", minute=20)
    ids["arch"] = add("arch", "Alpha Shelf", status="archived", minute=21)
    ids["marg"] = add("marg", "Beta Shelf", kv="marginal", minute=22)

    run = fx.sql(
        "INSERT INTO recluster_runs (user_id, status, completed_at) VALUES (%s,'completed',NOW())"
        " RETURNING id",
        (user["id"],),
    )[0][0]
    cl = fx.sql(
        "INSERT INTO clusters (user_id, cluster_slug, cluster_name, recluster_run)"
        " VALUES (%s,'s','Old cluster name',%s) RETURNING id",
        (user["id"], run),
    )[0][0]
    for k in ("a0", "a1", "b0"):
        fx.sql("INSERT INTO page_clusters (page_id, cluster_id) VALUES (%s,%s)", (ids[k][0], cl))

    from backend.services import text_embedder

    monkeypatch.setattr(
        text_embedder,
        "embed_texts",
        lambda texts, user_id=None: np.array(
            [_vec(0) if t.startswith("alpha") else _vec(90) for t in texts]
        ),
    )
    return {"user": user["id"], "ids": ids}


def _read_csv(p):
    with open(p, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


@db
def test_fetch_candidates_active_labelled_with_cluster(world):
    cs = {c.content_id: c for c in sa.fetch_candidates(world["user"])}
    ids = world["ids"]
    assert ids["arch"][1] not in cs  # archived page left out
    assert cs[ids["a0"][1]].cluster_name == "Old cluster name"
    assert cs[ids["a2"][1]].cluster_name is None
    assert cs[ids["a0"][1]].url.startswith("https://alpha0.example.test")
    assert len(cs) == 12  # 13 labelled content rows, the archived one left out


@db
def test_dry_run_writes_files_but_nothing_to_the_database(world, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(
        "backend.services.llm_service.LLMService",
        lambda: (_ for _ in ()).throw(AssertionError("model used in dry run")),
    )
    out = tmp_path / "sample.csv"
    assert (
        sa.main(["--user-id", str(world["user"]), "--out", str(out), "--n", "6", "--seed", "1"])
        == 0
    )
    printed = capsys.readouterr().out
    assert "Dry run" in printed and "Sample                : 6 pages" in printed
    assert "Synthetic" not in printed and "example.test" not in printed  # counts only
    for table in ("shelves", "shelf_pages", "shelf_strings", "cost_events"):
        assert fx.sql(f"SELECT COUNT(*) FROM {table}")[0][0] == 0
    rows = _read_csv(out)
    key = _read_csv(tmp_path / "sample.key.csv")
    assert len(rows) == len(key) == 6
    assert {r["mode"] for r in key} == {"dry-run"}
    exclude_cid = world["ids"]["x"][1]
    assert exclude_cid not in {int(r["page_content_id"]) for r in key}  # excluded never drawn
    assert world["ids"]["marg"][1] not in {int(r["page_content_id"]) for r in key}  # nor marginal
    assert {r["shelf_name"] for r in key} <= {"Alpha Shelf", "Beta Shelf"}
    assert settings.label_first_enabled is False


@db
def test_existing_files_are_not_overwritten(world, tmp_path, capsys):
    out = tmp_path / "sample.csv"
    out.write_text("owner marks")
    assert sa.main(["--user-id", str(world["user"]), "--out", str(out)]) == 2
    assert out.read_text() == "owner marks"
    assert "--overwrite" in capsys.readouterr().err
    assert (
        sa.main(["--user-id", str(world["user"]), "--out", str(out), "--overwrite", "--n", "3"])
        == 0
    )
    assert out.read_text() != "owner marks"


@db
def test_apply_without_tty_or_yes_writes_nothing(world, tmp_path, monkeypatch, capsys):
    monkeypatch.setattr("sys.stdin.isatty", lambda: False, raising=False)
    out = tmp_path / "sample.csv"
    assert sa.main(["--user-id", str(world["user"]), "--out", str(out), "--apply"]) == 1
    assert "--yes" in capsys.readouterr().err
    assert fx.sql("SELECT COUNT(*) FROM shelves")[0][0] == 0 and not out.exists()


@db
def test_apply_persists_shelves_with_the_flag_still_off(world, tmp_path, capsys):
    out = tmp_path / "sample.xlsx"
    args = ["--user-id", str(world["user"]), "--out", str(out), "--apply", "--yes", "--n", "8"]
    assert settings.label_first_enabled is False
    assert sa.main(args) == 0
    shelves = fx.sql("SELECT name, n_pages FROM shelves ORDER BY name")
    assert shelves == [
        ("Alpha Shelf", 6),
        ("Beta Shelf", 5),
    ]  # archived + excluded left out; marginal shelved
    assert fx.sql("SELECT COUNT(*) FROM shelf_pages")[0][0] == 11
    _, key = sa.read_table(tmp_path / "sample.key.csv")
    assert len(key) == 8 and {k[7] for k in key} == {"applied"}
    # re-running assigns nothing new and does not change the shelves
    out2 = tmp_path / "again.csv"
    assert (
        sa.main(
            ["--user-id", str(world["user"]), "--out", str(out2), "--apply", "--yes", "--n", "8"]
        )
        == 0
    )
    assert fx.sql("SELECT name, n_pages FROM shelves ORDER BY name") == shelves


@db
def test_apply_maintain_runs_one_pass_without_the_flag(world, tmp_path, monkeypatch):
    calls = []
    from backend.services import shelf_service

    real = shelf_service.weekly_maintenance
    monkeypatch.setattr(
        shelf_service,
        "weekly_maintenance",
        lambda uid, **kw: (
            calls.append(kw)
            or real(uid, ask_merge=lambda a, b: False, force=kw.get("force", False))
        ),
    )
    out = tmp_path / "m.csv"
    args = [
        "--user-id",
        str(world["user"]),
        "--out",
        str(out),
        "--apply",
        "--yes",
        "--maintain",
        "--n",
        "4",
    ]
    assert sa.main(args) == 0
    assert calls == [{"force": True}]


@db
def test_no_labels_is_a_clear_error(tmp_path, capsys):
    fx.sql(
        "TRUNCATE shelf_strings, shelf_pages, shelves, page_labels, cost_events, graph_cache,"
        " cluster_edges, page_clusters, clusters, recluster_runs, pages, page_content, captures,"
        " users CASCADE"
    )
    user = fx.make_user("empty")
    assert sa.main(["--user-id", str(user["id"]), "--out", str(tmp_path / "s.csv")]) == 2
    assert "label backfill" in capsys.readouterr().err


@db
def test_flag_off_shelf_entry_points_still_do_nothing(world):
    from backend.services import shelf_service

    assert shelf_service.assign_all_unshelved(world["user"]) == {}
    assert shelf_service.weekly_maintenance(world["user"]) == {}
    assert fx.sql("SELECT COUNT(*) FROM shelves")[0][0] == 0
